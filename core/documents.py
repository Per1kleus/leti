"""Reading what is in a file without putting the file in the prompt.

Leti could already read a text file and load a dataset. What it could not do was
the ordinary thing a person means by "read these PDFs and compare the offers":
look at a folder, work out which files the question is actually about, open them,
and quote the part that answers it - with somewhere to point afterwards.

The whole design is about NOT loading everything. A folder of ten contracts is
several hundred thousand tokens; the answer to "which one is cheaper" is two
numbers. So a request goes through four stages, each one cheaper than reading:

    find_relevant   which files could this be about, by name and kind and where
                    they live. No contents are opened at all.
    describe        what a file is: type, size, how many pages or sheets or rows.
                    Structure, not text.
    outline         the shape of one file - its pages, sheets or headings, with a
                    line from each. Enough to choose where to read.
    extract         the parts that answer the question, and only those, each one
                    carrying the label it came from.

That last part is what makes a citation possible. Every piece of text this module
returns is attached to a SECTION - "Page 7", "Sheet 'Pricing'", "Rows 51-100",
"## Payment terms" - because a quotation nobody can locate is worth less than no
quotation, and a located one Leti can be held to. Nothing here invents a label:
a section's name comes from the file's own structure.

Selection is deterministic: term overlap against the file's name and its sections,
scored the way core/tool_router.py scores tools. There is no second model deciding
what is relevant, and no embedding index to keep in step with the disk.

What it cannot read, it says so about. A corrupt PDF, a format with no reader
installed, a file that is not there - each comes back as a plain refusal with a
reason, never as an empty extraction that reads like an empty document.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import math
import re
import time
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

logger = logging.getLogger("leti.documents")

# What a section is allowed to carry, and how much of a file can come back from
# one extraction. Both are budgets on the PROMPT, not on the file: a 400-page PDF
# is fine, it just cannot arrive all at once.
MAX_SECTION_CHARS = 4_000
DEFAULT_EXTRACT_CHARS = 6_000
MAX_EXTRACT_CHARS = 24_000
MAX_SECTIONS_RETURNED = 12
OUTLINE_PREVIEW_CHARS = 140
MAX_FILE_BYTES = 80 * 1024 * 1024      # past this, reading is a disk problem
# An outline is meant to be structure, not contents. Sixty sections at 140
# characters each is 8KB of "preview" - which is contents by another name.
MAX_OUTLINE_SECTIONS = 40
MAX_OUTLINE_CHARS = 3_000
# How much of one comparison may reach the model in total, however many files
# were named. Per-file budgets alone multiply.
MAX_COMPARE_CHARS = 12_000
# How many pages of a PDF are opened to answer one question. Extracting text
# from four hundred pages to find one is minutes of work and hundreds of
# megabytes, so a long document is scanned in a window and says so - and a
# specific page can always be asked for by name.
MAX_PDF_PAGES_SCANNED = 120
MAX_DOCX_TABLE_ROWS = 40               # per table section; more are chunked

# Extensions Leti can open, and what it does with each.
KINDS: Dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text", ".log": "text", ".rst": "text",
    ".md": "markdown", ".markdown": "markdown",
    ".csv": "csv", ".tsv": "csv",
    ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".json": "json",
    ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".gif": "image", ".bmp": "image", ".webp": "image",
}

# Kinds whose bytes are not their text: decoding one as UTF-8 produces noise, so
# read_file sends these through this module rather than through read_text().
BINARY_KINDS = frozenset({"pdf", "docx", "xlsx"})

_WORD = re.compile(r"[a-z0-9]+")
# Words that match everything and therefore rank nothing.
_STOPWORDS = frozenset("""
a an and are as at be been but by can could do does for from has have how in into
is it its me my of on or our so than that the their them then there these they
this those to us was were what when where which who why will with you your
file files document documents please tell show find read give get compare
""".split())


# --------------------------------------------------------------------------- #
# What a file is
# --------------------------------------------------------------------------- #

def kind_of(path: Path) -> Optional[str]:
    return KINDS.get(path.suffix.lower())


def is_supported(path: Path) -> bool:
    return kind_of(path) is not None


def describe(path: Any) -> Dict[str, Any]:
    """Type, size and shape. Opens the file only far enough to count its parts."""
    p = Path(path)
    if not p.exists():
        return _problem(p, "there is no file at that path")
    if not p.is_file():
        return _problem(p, "that is a folder, not a file")

    stat = p.stat()
    info: Dict[str, Any] = {
        "ok": True,
        "path": str(p),
        "name": p.name,
        "kind": kind_of(p) or "unsupported",
        "bytes": stat.st_size,
        "modified": time.strftime("%Y-%m-%d %H:%M", time.localtime(stat.st_mtime)),
    }
    if info["kind"] == "unsupported":
        info["ok"] = False
        info["error"] = (f"Leti has no reader for '{p.suffix or 'a file with no extension'}'. "
                         f"Readable kinds: {', '.join(sorted(set(KINDS.values())))}.")
        return info
    if stat.st_size > MAX_FILE_BYTES:
        info["ok"] = False
        info["error"] = f"{stat.st_size:,} bytes is past the {MAX_FILE_BYTES:,} byte limit."
        return info

    try:
        info.update(_shape(p, info["kind"]))
    except Exception as e:
        info["ok"] = False
        info["error"] = f"Couldn't read the structure of {p.name}: {e}"
    return info


def _problem(path: Path, why: str) -> Dict[str, Any]:
    return {"ok": False, "path": str(path), "name": path.name, "error": why}


def _shape(path: Path, kind: str) -> Dict[str, Any]:
    """How many parts the file has, without reading all of them."""
    if kind == "pdf":
        reader = _pdf_reader(path)
        return {"pages": len(reader.pages)}
    if kind == "docx":
        paragraphs, tables = _docx_body(path)
        return {"paragraphs": len(paragraphs), "tables": tables}
    if kind == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return {"sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column}
                               for s in book.worksheets]}
        finally:
            book.close()
    if kind == "csv":
        with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.reader(fh, delimiter=_csv_delimiter(path))
            header = next(reader, [])
            rows = sum(1 for _ in reader)
        return {"columns": header, "rows": rows}
    if kind == "json":
        data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        if isinstance(data, dict):
            return {"top_level_keys": sorted(data)[:60]}
        if isinstance(data, list):
            return {"items": len(data)}
        return {"json_type": type(data).__name__}
    if kind == "image":
        # Three different things get confused with each other here, so each is
        # named separately and one of them is honestly absent:
        #   metadata          - the file's own properties. Available, below.
        #   image understanding - the vision model saying what a picture shows.
        #                      Available, through look_at_image.
        #   text extraction (OCR) - reading characters off the pixels. NOT
        #                      available: nothing in this project does OCR, and
        #                      a vision model's guess at a serial number is not
        #                      the same thing as having read it.
        shape: Dict[str, Any] = {
            "text_extraction": "unavailable",
            "text_extraction_note": (
                "Leti has no OCR. Text visible in this image cannot be extracted. "
                "look_at_image can describe the picture with the vision model, which "
                "may transcribe some text but can misread it - report anything it says "
                "about text as the model's reading of an image, never as the file's "
                "contents."),
            "image_understanding": "look_at_image",
        }
        try:
            from PIL import Image

            with Image.open(path) as im:
                shape.update({"width": im.width, "height": im.height, "format": im.format,
                              "mode": im.mode})
        except Exception as e:
            shape["note"] = f"An image Leti could not open ({e})."
        return shape
    text = path.read_text(encoding="utf-8", errors="replace")
    return {"characters": len(text), "lines": text.count("\n") + 1}


# --------------------------------------------------------------------------- #
# Sections: the unit of citation
# --------------------------------------------------------------------------- #

def sections_and_scan(path: Any, first_page: int = 1) -> Tuple[List[Dict[str, str]],
                                                                 Dict[str, Any]]:
    """The file in labelled pieces, and what was opened to produce them.

    The second half matters for a long PDF: only a window of pages is read, so
    the answer has to say which ones, or "not in this document" would really mean
    "not in the first hundred pages" and nobody could tell the difference.
    """
    p = Path(path)
    kind = kind_of(p)
    if kind is None:
        raise ValueError(f"No reader for '{p.suffix}'.")
    scan: Dict[str, Any] = {}

    if kind == "pdf":
        total, window = pdf_window(p, first_page)
        scan = {"pages": total, "pages_scanned": window, "first_page_scanned": first_page}
        raw = list(_pdf_sections(p, first_page=first_page))
        # Pages with nothing on them are the signature of a scan. Counting them
        # is what lets extract() say "this is images of text" rather than
        # "nothing matched", which are different problems with different answers.
        scan["pages_without_text"] = sum(1 for _, text in raw if not (text or "").strip())
    else:
        reader = {
            "docx": _docx_sections, "xlsx": _xlsx_sections, "csv": _csv_sections,
            "json": _json_sections, "markdown": _markdown_sections,
            "text": _text_sections, "image": _image_sections,
        }[kind]
        raw = list(reader(p))

    out = []
    for label, text in raw:
        text = (text or "").strip()
        if not text:
            continue
        out.append({"label": label, "text": text[:MAX_SECTION_CHARS]})
    return out, scan


def sections(path: Any) -> List[Dict[str, str]]:
    """The file in labelled pieces. Labels come from the file, never invented."""
    return sections_and_scan(path)[0]


def _pdf_reader(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError as e:      # pragma: no cover - depends on the install
        raise RuntimeError(
            "Reading PDFs needs the 'pypdf' package (pip install pypdf). Leti will not "
            "guess at a PDF's contents without it."
        ) from e
    # strict=False: a PDF with a damaged cross-reference table or an off-by-one
    # object offset is usually still readable, and refusing the whole file over a
    # structural complaint is a worse answer than reading what is there.
    return PdfReader(str(path), strict=False)


def pdf_window(path: Path, first: int = 1, limit: int = MAX_PDF_PAGES_SCANNED) -> Tuple[int, int]:
    """How many pages this PDF has, and how many of them one pass will open."""
    reader = _pdf_reader(path)
    total = len(reader.pages)
    return total, min(total, max(0, total - (first - 1)), limit)


def _pdf_sections(path: Path, first_page: int = 1,
                  limit: int = MAX_PDF_PAGES_SCANNED) -> Iterable[Tuple[str, str]]:
    """Pages, one section each, from `first_page`, at most `limit` of them.

    A page that will not extract yields empty rather than stopping the document:
    one damaged page out of two hundred should cost that page, not the file.
    """
    reader = _pdf_reader(path)
    pages = reader.pages
    start = max(0, first_page - 1)
    for number in range(start, min(len(pages), start + max(0, limit))):
        try:
            text = pages[number].extract_text() or ""
        except Exception as e:
            logger.debug(f"{path.name} page {number + 1} would not extract: {e}")
            text = ""
        yield f"Page {number + 1}", text


# --- DOCX. A .docx is a zip of XML, so this needs no package: the alternative
#     was a dependency to read a format the standard library can already open.
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_blocks(path: Path) -> List[Dict[str, Any]]:
    """The document body in order: paragraphs and tables, each appearing once.

    The previous version walked every descendant of the body, which meant the
    paragraphs INSIDE a table came back as loose paragraphs - the cells were
    there, stripped of the thing that made them a table, and the table itself was
    only a count. Walking children and stopping at a table fixes both halves: the
    rows keep their shape, and their text is not also reported as prose.
    """
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find(f"{_W}body")
    if body is None:
        return []

    def text_of(node) -> str:
        return "".join(t.text or "" for t in node.iter(f"{_W}t"))

    def rows_of(table) -> List[List[str]]:
        rows = []
        for row in table.findall(f"{_W}tr"):
            cells = [" ".join(text_of(cell).split())
                     for cell in row.findall(f"{_W}tc")]
            if any(c for c in cells):
                rows.append(cells)
        return rows

    blocks: List[Dict[str, Any]] = []

    def walk(parent) -> None:
        for node in parent:
            if node.tag == f"{_W}p":
                style_node = node.find(f"{_W}pPr/{_W}pStyle")
                style = style_node.get(f"{_W}val", "") if style_node is not None else ""
                text = text_of(node)
                if text.strip():
                    blocks.append({"kind": "paragraph", "style": style, "text": text})
            elif node.tag == f"{_W}tbl":
                rows = rows_of(node)
                if rows:
                    blocks.append({"kind": "table", "rows": rows})
            elif node.tag in (f"{_W}sdt", f"{_W}sdtContent", f"{_W}smartTag"):
                # Content controls wrap real paragraphs; a table can live in one.
                walk(node)

    walk(body)
    return blocks


def _docx_body(path: Path) -> Tuple[List[Tuple[str, str]], int]:
    """Paragraphs as (style, text) and the number of tables. Kept for describe()."""
    blocks = _docx_blocks(path)
    paragraphs = [(b["style"], b["text"]) for b in blocks if b["kind"] == "paragraph"]
    return paragraphs, sum(1 for b in blocks if b["kind"] == "table")


def _docx_sections(path: Path) -> Iterable[Tuple[str, str]]:
    """Headings split the prose; each table is its own section, where it stands.

    A table becomes tab-separated rows with its header repeated when it is long
    enough to be chunked - the same shape a spreadsheet or a CSV comes back as,
    so "the pricing table in Contract_B.docx" reads the same whichever file it
    was in. Its label says where it was, so a figure taken from it can be
    attributed to a table rather than to the document in general.
    """
    blocks = _docx_blocks(path)
    label, buffer, number, table_number = None, [], 0, 0

    def flush():
        nonlocal buffer, number, label
        if not buffer:
            return None
        number += 1
        out = (label or f"Section {number}", "\n".join(buffer))
        buffer = []
        return out

    for block in blocks:
        if block["kind"] == "table":
            pending = flush()
            if pending:
                yield pending
            table_number += 1
            heading = f" under '{label}'" if label else ""
            rows = block["rows"]
            header = "\t".join(rows[0]) if rows else ""
            for start in range(0, len(rows), MAX_DOCX_TABLE_ROWS):
                chunk = rows[start:start + MAX_DOCX_TABLE_ROWS]
                lines = ["\t".join(cells) for cells in chunk]
                if start and header:
                    lines.insert(0, header)      # a cited chunk still reads
                part = (f" rows {start + 1}-{start + len(chunk)}"
                        if len(rows) > MAX_DOCX_TABLE_ROWS else "")
                yield f"Table {table_number}{heading}{part}", "\n".join(lines)
            continue

        style, text = block["style"], block["text"]
        if style.lower().startswith("heading"):
            pending = flush()
            if pending:
                yield pending
            label = text.strip()[:80]
            continue
        buffer.append(text)
        if sum(len(b) for b in buffer) > MAX_SECTION_CHARS:
            pending = flush()
            if pending:
                yield pending
            label = f"{label} (continued)" if label else None
    pending = flush()
    if pending:
        yield pending


def _xlsx_sections(path: Path) -> Iterable[Tuple[str, str]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            lines, header = [], None
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if header is None:
                    header = cells
                lines.append("\t".join(cells))
                if sum(len(line) for line in lines) > MAX_SECTION_CHARS:
                    break
            yield f"Sheet '{sheet.title}'", "\n".join(lines)
    finally:
        book.close()


def _csv_delimiter(path: Path) -> str:
    return "\t" if path.suffix.lower() == ".tsv" else ","


def _csv_sections(path: Path) -> Iterable[Tuple[str, str]]:
    """Row ranges, with the header repeated so a cited chunk still reads."""
    with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=_csv_delimiter(path)))
    if not rows:
        return
    header, body = rows[0], rows[1:]
    header_line = "\t".join(header)
    chunk, start = [], 1
    for index, row in enumerate(body, start=1):
        chunk.append("\t".join(str(c) for c in row))
        if sum(len(c) for c in chunk) > MAX_SECTION_CHARS:
            yield f"Rows {start}-{index}", header_line + "\n" + "\n".join(chunk)
            chunk, start = [], index + 1
    if chunk:
        yield f"Rows {start}-{len(body)}", header_line + "\n" + "\n".join(chunk)


def _json_sections(path: Path) -> Iterable[Tuple[str, str]]:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    if isinstance(data, dict):
        for key in sorted(data):
            yield f"key '{key}'", json.dumps(data[key], indent=1, default=str)
        return
    if isinstance(data, list):
        for start in range(0, len(data), 25):
            block = data[start:start + 25]
            yield (f"items {start + 1}-{start + len(block)}",
                   json.dumps(block, indent=1, default=str))
        return
    yield "value", json.dumps(data, default=str)


def _markdown_sections(path: Path) -> Iterable[Tuple[str, str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    label, buffer = None, []
    for line in text.splitlines():
        if line.startswith("#"):
            if buffer:
                yield label or "Before the first heading", "\n".join(buffer)
                buffer = []
            label = line.strip()[:80]
            continue
        buffer.append(line)
    if buffer:
        yield label or "Whole file", "\n".join(buffer)


def _text_sections(path: Path) -> Iterable[Tuple[str, str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    chunk, start = [], 1
    for index, line in enumerate(lines, start=1):
        chunk.append(line)
        if sum(len(c) for c in chunk) > MAX_SECTION_CHARS:
            yield f"Lines {start}-{index}", "\n".join(chunk)
            chunk, start = [], index + 1
    if chunk:
        yield f"Lines {start}-{len(lines)}", "\n".join(chunk)


def _image_sections(path: Path) -> Iterable[Tuple[str, str]]:
    """An image has no sections of text. What it has is a description of itself,
    and a plain statement that its text cannot be read."""
    shape = describe(path)
    yield "image", (f"{path.name}: {shape.get('width', '?')}x{shape.get('height', '?')} "
                    f"{shape.get('format', '')}. There is no extractable text in an image "
                    "and Leti has no OCR. Use look_at_image to have the vision model "
                    "describe what it shows; do not report that description as the file's "
                    "text.")


# --------------------------------------------------------------------------- #
# Outline and extraction
# --------------------------------------------------------------------------- #

def outline(path: Any) -> Dict[str, Any]:
    """The shape of a file: its parts, and a line from each. Enough to choose."""
    p = Path(path)
    info = describe(p)
    if not info.get("ok"):
        return info
    try:
        parts, scan = sections_and_scan(p)
    except Exception as e:
        return {**info, "ok": False, "error": f"Couldn't open {p.name}: {e}"}
    info.update(scan)
    info["sections_total"] = len(parts)

    # An outline is structure. A preview per section is what makes it usable, and
    # forty of them at 140 characters is where that stops being structure and
    # starts being the document, so both ends are capped and the cut is stated.
    shown, used = [], 0
    for part in parts[:MAX_OUTLINE_SECTIONS]:
        preview = " ".join(part["text"].split())[:OUTLINE_PREVIEW_CHARS]
        if used + len(preview) > MAX_OUTLINE_CHARS:
            break
        shown.append({"label": part["label"], "preview": preview})
        used += len(preview)
    info["sections"] = shown
    if len(shown) < len(parts):
        info["sections_not_shown"] = len(parts) - len(shown)
    if not parts:
        info["ok"] = False
        info["error"] = _nothing_readable(p, info)
    return info


def _nothing_readable(path: Path, info: Dict[str, Any]) -> str:
    """Why a file that opened produced no text. Said plainly, never guessed past."""
    if info.get("kind") == "pdf" and info.get("pages"):
        return (f"{path.name} has {info['pages']} page(s) and no extractable text in the "
                f"{info.get('pages_scanned', 0)} scanned. That is what a scanned document "
                "looks like: the pages are images. Leti has no OCR, so there is no text to "
                "read out of it - say so rather than describing what it might contain.")
    return (f"{path.name} opened, but there is no readable text in it. It may be empty, or "
            "its contents may be images rather than text.")


def _stem(word: str) -> str:
    """Just enough to make "offers" match Offer_Company_A.pdf.

    Deliberately crude and deliberately visible: a real stemmer would be a
    dependency and a source of surprises, and the whole job here is matching a
    plural in a question against a singular in a filename.
    """
    if len(word) > 4 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 4 and word.endswith(("ses", "xes", "zes", "ches", "shes")):
        return word[:-2]
    if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


_RANGE = re.compile(r"^\s*(?P<what>[a-z ']*?)\s*(?P<from>\d+)\s*[-\u2013to]+\s*(?P<to>\d+)\s*$",
                    re.I)
_NUMBERED = re.compile(r"(\d+)")


def _first_page_of(section: str) -> int:
    """Which page a request names, so a long PDF can be opened there.

    "Page 340" in a four-hundred-page document is past the scan window, and
    reading from page one to reach it would be the expensive thing this avoids.
    """
    text = str(section or "")
    if "page" not in text.lower():
        return 1
    found = _NUMBERED.search(text)
    return max(1, int(found.group(1))) if found else 1


def _sections_matching(parts: List[Dict[str, str]], section: str) -> List[int]:
    """Sections a request names - one label, or a range like 'Page 3-7'."""
    wanted = str(section or "").strip()
    span = _RANGE.match(wanted)
    if span:
        low, high = sorted((int(span.group("from")), int(span.group("to"))))
        prefix = span.group("what").strip().lower()
        chosen = []
        for index, part in enumerate(parts):
            label = part["label"].lower()
            if prefix and prefix not in label:
                continue
            numbers = [int(n) for n in _NUMBERED.findall(part["label"])]
            if numbers and low <= numbers[0] <= high:
                chosen.append(index)
        if chosen:
            return chosen
    return [i for i, s in enumerate(parts) if wanted.lower() in s["label"].lower()]


def _terms(text: str) -> List[str]:
    return [_stem(w) for w in _WORD.findall(str(text or "").lower())
            if len(w) > 2 and w not in _STOPWORDS]


def _score_sections(parts: List[Dict[str, str]], query: str) -> List[Tuple[float, int]]:
    """Rank sections against the question. Rare words count for more.

    The same shape of scoring core/tool_router.py uses, for the same reason: a
    word that appears in every section separates nothing.
    """
    wanted = set(_terms(query))
    if not wanted:
        return [(0.0, i) for i in range(len(parts))]
    haystacks = [set(_terms(p["label"] + " " + p["text"])) for p in parts]
    total = len(parts) or 1
    scored = []
    for index, words in enumerate(haystacks):
        score = 0.0
        for term in wanted & words:
            appearing = sum(1 for h in haystacks if term in h) or 1
            score += math.log(1 + total / appearing)
        scored.append((score, index))
    scored.sort(key=lambda pair: (-pair[0], pair[1]))
    return scored


def extract(path: Any, query: str = "", max_chars: int = DEFAULT_EXTRACT_CHARS,
            section: str = "") -> Dict[str, Any]:
    """The parts of one file that answer the question, each carrying its label.

    With no query this is the beginning of the file rather than all of it, which
    is the honest answer to "what is in here" - and `sections_total` says how much
    was left, so asking for more is a decision rather than a guess.
    """
    p = Path(path)
    info = describe(p)
    if not info.get("ok"):
        return info
    first_page = _first_page_of(section) if section else 1
    try:
        parts, scan = sections_and_scan(p, first_page=first_page)
    except Exception as e:
        return {**info, "ok": False, "error": f"Couldn't read {p.name}: {e}"}
    info.update(scan)
    if not parts:
        # A page number past the end is a different problem from a scan, and
        # saying "this looks like images of text" about a file that simply has
        # fewer pages than you asked for would be a confident wrong answer.
        if section and info.get("pages") and first_page > info["pages"]:
            return {**info, "ok": False,
                    "error": f"{p.name} has {info['pages']} page(s); '{section}' is past "
                             "the end of it."}
        return {**info, "ok": False, "error": _nothing_readable(p, info)}

    budget = max(500, min(int(max_chars or DEFAULT_EXTRACT_CHARS), MAX_EXTRACT_CHARS))

    if section:
        wanted = _sections_matching(parts, section)
        if not wanted:
            return {**info, "ok": False,
                    "error": f"{p.name} has no section matching '{section}'.",
                    "available": [s["label"] for s in parts[:40]]}
        chosen = wanted
    else:
        ranked = _score_sections(parts, query)
        chosen = []
        used = 0
        for score, index in ranked:
            if len(chosen) >= MAX_SECTIONS_RETURNED:
                break
            if query and score <= 0 and chosen:
                break            # nothing more in this file speaks to the question
            length = len(parts[index]["text"])
            if used and used + length > budget:
                continue
            chosen.append(index)
            used += length
        if not chosen:
            chosen = [ranked[0][1]]

    chosen.sort()
    returned, used = [], 0
    result: Dict[str, Any] = {}
    for index in chosen:
        part = parts[index]
        room = budget - used
        if room <= 0:
            break
        text = part["text"][:room]
        returned.append({"label": part["label"], "text": text,
                         "truncated": len(text) < len(part["text"])})
        used += len(text)

    result = {
        **info,
        "source": p.name,
        "query": query or None,
        "sections": returned,
        "sections_total": len(parts),
        "sections_returned": len(returned),
        "characters": used,
        "more_available": len(returned) < len(parts),
        "how_to_cite": (f"Attribute anything quoted here to {p.name} and the section "
                        "label it came from. Do not cite a label that is not listed."),
    }

    # A long PDF was read in a window. Saying so is the difference between "not in
    # this document" and "not in the hundred pages that were opened".
    if info.get("pages") and info.get("pages_scanned", 0) < info["pages"]:
        last = info["first_page_scanned"] + info["pages_scanned"] - 1
        result["note"] = (
            f"Pages {info['first_page_scanned']}-{last} of {info['pages']} were read. "
            "Ask for a later page by name (section='Page 200') to read further; nothing "
            "here says anything about the pages that were not opened.")
    elif info.get("pages_without_text"):
        result["note"] = (
            f"{info['pages_without_text']} of the {info['pages_scanned']} pages read had "
            "no extractable text - those pages are images. Leti has no OCR, so nothing is "
            "known about what is on them.")
    return result


# --------------------------------------------------------------------------- #
# Which files a request is about
# --------------------------------------------------------------------------- #

def candidate_files(folder: Any, recursive: bool = True) -> List[Path]:
    root = Path(folder)
    if not root.is_dir():
        return []
    walk = root.rglob("*") if recursive else root.glob("*")
    found = []
    try:
        for path in walk:
            try:
                if path.is_file() and is_supported(path):
                    found.append(path)
            except OSError:
                continue          # a broken link or a file that vanished mid-walk
    except OSError as e:
        logger.warning(f"Couldn't finish listing {root}: {e}")
    return found


def find_relevant(request: str, files: Iterable[Any], limit: int = 12,
                  look_inside: int = 0) -> List[Dict[str, Any]]:
    """Rank files against a request by what can be known cheaply.

    Names, folders and kinds first, because they cost nothing and are usually
    enough - "Offer_Company_A.pdf" answers "compare the offers" on its name
    alone. `look_inside` opens at most that many of the best remaining candidates
    to check, which is the expensive path and therefore never the first one.
    """
    wanted = set(_terms(request))
    paths = [Path(f) for f in files]
    scored: List[Dict[str, Any]] = []
    for path in paths:
        name_words = set(_terms(str(path)))
        hits = sorted(wanted & name_words)
        try:
            size = path.stat().st_size
        except OSError:
            # Listed a moment ago and gone now, or unreadable. Rank it anyway and
            # let opening it be the thing that reports the problem.
            size = 0
        scored.append({
            "path": str(path), "name": path.name, "kind": kind_of(path) or "unsupported",
            "bytes": size,
            "score": float(len(hits)),
            "matched": hits,
            "why": (f"the name matches {', '.join(hits)}" if hits else "same folder"),
        })

    scored.sort(key=lambda f: (-f["score"], f["name"]))
    if look_inside > 0 and wanted:
        for entry in scored[:look_inside]:
            if entry["score"] > 0:
                continue
            try:
                inside = _score_sections(sections(entry["path"]), request)
            except Exception:
                continue
            best = inside[0][0] if inside else 0.0
            if best > 0:
                entry["score"] = round(best, 3)
                entry["why"] = "the question's words appear inside it"
        scored.sort(key=lambda f: (-f["score"], f["name"]))

    # When nothing matched by name there is no relevance to report, only a folder
    # listing - which is the honest thing to hand back rather than a ranking.
    return scored[:limit]
